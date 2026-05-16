from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import json
import re
from collections import defaultdict

# Parse SwanLab backup file
filepath = '/home/mclab/mjp/unifyimmun/source/swanlog/run-20260512_184254-7upkoa2qzq7oixuwtilk6/backup.swanlab'

with open(filepath, 'rb') as f:
    content = f.read()

scalar_data = defaultdict(list)
log_data = []
experiment_info = {}

lines = content.split(b'\n')
for line in lines:
    try:
        text = line.decode('utf-8', errors='ignore')
        if '{' in text and '"model_type"' in text:
            start = text.find('{')
            brace_count = 0
            end = start
            for i, c in enumerate(text[start:], start):
                if c == '{':
                    brace_count += 1
                elif c == '}':
                    brace_count -= 1
                    if brace_count == 0:
                        end = i + 1
                        break

            json_str = text[start:end]
            data = json.loads(json_str)

            if data.get('model_type') == 'Scalar':
                scalar_info = data.get('data', {})
                key = scalar_info.get('key', 'unknown')
                metric = scalar_info.get('metric', {})
                value = metric.get('data')
                epoch = scalar_info.get('epoch', 0)
                if value is not None:
                    scalar_data[key].append({'epoch': epoch, 'value': value})
            elif data.get('model_type') == 'Log':
                log_data.append(data.get('data', {}))
            elif data.get('model_type') == 'Experiment':
                experiment_info = data.get('data', {})
    except Exception as e:
        pass

# Get final metrics
final_metrics = {}
for key in scalar_data:
    if not key.startswith('__swanlab__'):
        values = scalar_data[key]
        if values:
            sorted_values = sorted(values, key=lambda x: x['epoch'])
            final_metrics[key] = sorted_values[-1]['value']

# Create PPT
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

# Background color
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0xF5, 0xF5, 0xF5)

# Title
title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.3), Inches(12.7), Inches(0.7))
title_frame = title_box.text_frame
title_para = title_frame.paragraphs[0]
title_para.text = "TCR_ESM2_2 Training Results Analysis"
title_para.font.size = Pt(32)
title_para.font.bold = True
title_para.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
title_para.alignment = PP_ALIGN.CENTER

# Subtitle - Experiment info
subtitle_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.9), Inches(12.7), Inches(0.5))
subtitle_frame = subtitle_box.text_frame
subtitle_para = subtitle_frame.paragraphs[0]
subtitle_para.text = f"SwanLab Run: run-20260512_184254 | Best Epoch: {int(final_metrics.get('best_epoch', 7))} | Best Avg: {final_metrics.get('best_performance_avg', 0):.4f}"
subtitle_para.font.size = Pt(14)
subtitle_para.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
subtitle_para.alignment = PP_ALIGN.CENTER

# Section 1: Training Configuration
config_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(1.5), Inches(4), Inches(2.3))
config_shape.fill.solid()
config_shape.fill.fore_color.rgb = RGBColor(0x34, 0x49, 0x5E)
config_shape.line.fill.background()

config_title = config_shape.text_frame.paragraphs[0]
config_title.text = "Training Configuration"
config_title.font.size = Pt(16)
config_title.font.bold = True
config_title.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)

# Add config items
config_items = [
    ("Embedding", "ESM2_650M"),
    ("Batch Size", "64"),
    ("Epochs", "30 (Early Stop @12)"),
    ("Learning Rate", "1e-4"),
    ("d_model", "64"),
    ("Early Stop Patience", "5"),
]

for item_name, item_value in config_items:
    p = config_shape.text_frame.add_paragraph()
    p.text = f"{item_name}: {item_value}"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)
    p.space_before = Pt(4)

# Section 2: Training Metrics (Train & Val)
train_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(1.5), Inches(4), Inches(2.3))
train_shape.fill.solid()
train_shape.fill.fore_color.rgb = RGBColor(0x27, 0xAE, 0x60)
train_shape.line.fill.background()

train_title = train_shape.text_frame.paragraphs[0]
train_title.text = "Training & Validation Metrics"
train_title.font.size = Pt(16)
train_title.font.bold = True
train_title.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)

# Training metrics
train_metrics_to_show = [
    ("Train Loss", final_metrics.get('train_loss', 0)),
    ("Train AUC", final_metrics.get('train_roc_auc', 0)),
    ("Train Acc", final_metrics.get('train_accuracy', 0)),
    ("Train MCC", final_metrics.get('train_mcc', 0)),
    ("Val Loss", final_metrics.get('val_loss', 0)),
    ("Val AUC", final_metrics.get('val_roc_auc', 0)),
    ("Val Acc", final_metrics.get('val_accuracy', 0)),
    ("Val MCC", final_metrics.get('val_mcc', 0)),
]

for name, value in train_metrics_to_show:
    p = train_shape.text_frame.add_paragraph()
    p.text = f"{name}: {value:.4f}"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)
    p.space_before = Pt(2)

# Section 3: Test Set Performance
test_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.7), Inches(1.5), Inches(4.3), Inches(2.3))
test_shape.fill.solid()
test_shape.fill.fore_color.rgb = RGBColor(0xE7, 0x4C, 0x3C)
test_shape.line.fill.background()

test_title = test_shape.text_frame.paragraphs[0]
test_title.text = "Test Set Performance"
test_title.font.size = Pt(16)
test_title.font.bold = True
test_title.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)

# Test metrics
test_metrics = [
    ("Independent AUC", final_metrics.get('independent_roc_auc', 0)),
    ("Independent Acc", final_metrics.get('independent_accuracy', 0)),
    ("Triple AUC", final_metrics.get('triple_roc_auc', 0)),
    ("Triple Acc", final_metrics.get('triple_accuracy', 0)),
    ("Covid AUC", final_metrics.get('covid_roc_auc', 0)),
    ("Covid Acc", final_metrics.get('covid_accuracy', 0)),
]

for name, value in test_metrics:
    p = test_shape.text_frame.add_paragraph()
    p.text = f"{name}: {value:.4f}"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)
    p.space_before = Pt(4)

# Section 4: Detailed Metrics Table
table_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(4.0), Inches(6.3), Inches(3.0))
table_shape.fill.solid()
table_shape.fill.fore_color.rgb = RGBColor(0x34, 0x49, 0x5E)
table_shape.line.fill.background()

table_title = table_shape.text_frame.paragraphs[0]
table_title.text = "Complete Performance Metrics"
table_title.font.size = Pt(14)
table_title.font.bold = True
table_title.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)

# Add metrics table as text
metrics_lines = [
    "Metric         | Train   | Val     | Indep  | Triple | Covid",
    "─────────────────────────────────────────────────────────",
    f"ROC-AUC       | {final_metrics.get('train_roc_auc',0):.4f} | {final_metrics.get('val_roc_auc',0):.4f} | {final_metrics.get('independent_roc_auc',0):.4f} | {final_metrics.get('triple_roc_auc',0):.4f} | {final_metrics.get('covid_roc_auc',0):.4f}",
    f"Accuracy      | {final_metrics.get('train_accuracy',0):.4f} | {final_metrics.get('val_accuracy',0):.4f} | {final_metrics.get('independent_accuracy',0):.4f} | {final_metrics.get('triple_accuracy',0):.4f} | {final_metrics.get('covid_accuracy',0):.4f}",
    f"MCC           | {final_metrics.get('train_mcc',0):.4f} | {final_metrics.get('val_mcc',0):.4f} | 0.7298 | 0.5965 | 0.0793",
    f"F1 Score      | {final_metrics.get('train_f1',0):.4f} | {final_metrics.get('val_f1',0):.4f} | 0.8191 | 0.6882 | 0.2008",
    f"Sensitivity   | {final_metrics.get('train_sensitivity',0):.4f} | {final_metrics.get('val_sensitivity',0):.4f} | 0.7858 | 0.5517 | 0.1196",
    f"Specificity   | {final_metrics.get('train_specificity',0):.4f} | {final_metrics.get('val_specificity',0):.4f} | 0.9288 | 0.9670 | -",
]

for line in metrics_lines:
    p = table_shape.text_frame.add_paragraph()
    p.text = line
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)
    p.space_before = Pt(1)

# Section 5: Key Observations
obs_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.0), Inches(6.2), Inches(3.0))
obs_shape.fill.solid()
obs_shape.fill.fore_color.rgb = RGBColor(0xF3, 0x9C, 0x12)
obs_shape.line.fill.background()

obs_title = obs_shape.text_frame.paragraphs[0]
obs_title.text = "Key Observations"
obs_title.font.size = Pt(14)
obs_title.font.bold = True
obs_title.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)

observations = [
    "✓ Strong performance on Independent test set (AUC=0.9415)",
    "✓ Good generalization on Triple dataset (AUC=0.8349)",
    "✗ Poor performance on Covid dataset (AUC=0.5185)",
    "• Early stopping triggered at epoch 12",
    "• Best model saved at epoch 7",
    "• Training time: 37888 seconds (~10.5 hours)",
    "• FGM adversarial training applied for robustness",
    "• ESM2 encoder loaded from HLA_ESM2_2 (Phase 2)",
]

for obs in observations:
    p = obs_shape.text_frame.add_paragraph()
    p.text = obs
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(0xEC, 0xF0, 0xF1)
    p.space_before = Pt(4)

# Save PPT
output_path = '/home/mclab/mjp/unifyimmun/source/swanlog/TCR_ESM2_2_Analysis.pptx'
prs.save(output_path)
print(f"PPT saved to: {output_path}")