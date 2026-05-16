"""
Generate analysis PPT for Phase 2 (HLA_ESM2_2 and TCR_ESM2_2) training results.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


def add_phase2_slide(prs, title, data_dict, test_sets_data, comparison_data=None):
    """Add a Phase 2 analysis slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9.4), Inches(0.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)

    # Training summary
    summary_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.7), Inches(3), Inches(0.3))
    tf = summary_box.text_frame
    p = tf.paragraphs[0]
    p.text = "训练概况"
    p.font.size = Pt(14)
    p.font.bold = True

    # Summary table
    rows = 4
    cols = 2
    table = slide.shapes.add_table(rows, cols, Inches(0.3), Inches(1), Inches(3), Inches(1.2)).table

    summary_data = [
        ["项目", "值"],
        ["最佳 Epoch", data_dict["best_epoch"]],
        ["最佳验证性能均值", data_dict["best_performance_avg"]],
        ["早停计数", data_dict["no_improve_count"]],
    ]

    for i, row_data in enumerate(summary_data):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = cell_text
            cell.text_frame.paragraphs[0].font.size = Pt(11)
            if i == 0:
                cell.text_frame.paragraphs[0].font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(180, 180, 180)

    # Phase comparison table
    if comparison_data:
        comp_box = slide.shapes.add_textbox(Inches(0.3), Inches(2.3), Inches(4.5), Inches(0.3))
        tf = comp_box.text_frame
        p = tf.paragraphs[0]
        p.text = "Phase 1 vs Phase 2 对比"
        p.font.size = Pt(14)
        p.font.bold = True

        comp_rows = len(comparison_data) + 1
        comp_cols = 3
        comp_table = slide.shapes.add_table(comp_rows, comp_cols, Inches(0.3), Inches(2.6), Inches(4.5), Inches(comp_rows * 0.4)).table

        comp_header = ["对比项", "Phase 1", "Phase 2"]
        for j, header in enumerate(comp_header):
            cell = comp_table.cell(0, j)
            cell.text = header
            cell.text_frame.paragraphs[0].font.size = Pt(10)
            cell.text_frame.paragraphs[0].font.bold = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(180, 180, 180)

        for i, row_data in enumerate(comparison_data):
            for j, cell_text in enumerate(row_data):
                cell = comp_table.cell(i + 1, j)
                cell.text = str(cell_text)
                cell.text_frame.paragraphs[0].font.size = Pt(10)
                # Highlight improvements
                if j == 2 and cell_text.startswith("+"):
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(200, 255, 200)
                elif j == 2 and cell_text.startswith("-"):
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(255, 200, 200)

    # Test sets performance (right side)
    test_box = slide.shapes.add_textbox(Inches(5.3), Inches(0.7), Inches(4.2), Inches(0.3))
    tf = test_box.text_frame
    p = tf.paragraphs[0]
    p.text = "测试集表现"
    p.font.size = Pt(14)
    p.font.bold = True

    test_rows = len(test_sets_data) + 1
    test_cols = 3
    test_table = slide.shapes.add_table(test_rows, test_cols, Inches(5.3), Inches(1), Inches(4.2), Inches(test_rows * 0.45)).table

    test_header = ["测试集", "ROC-AUC", "Accuracy"]
    for j, header in enumerate(test_header):
        cell = test_table.cell(0, j)
        cell.text = header
        cell.text_frame.paragraphs[0].font.size = Pt(11)
        cell.text_frame.paragraphs[0].font.bold = True
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(180, 180, 180)

    for i, row_data in enumerate(test_sets_data):
        for j, cell_text in enumerate(row_data):
            cell = test_table.cell(i + 1, j)
            cell.text = str(cell_text)
            cell.text_frame.paragraphs[0].font.size = Pt(11)
            # Color coding
            if j == 1 or j == 2:
                try:
                    val = float(cell_text)
                    if val < 0.6:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor(255, 180, 180)  # Red for poor
                    elif val > 0.94:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor(180, 255, 180)  # Green for excellent
                except:
                    pass

    # Key observations
    obs_box = slide.shapes.add_textbox(Inches(5.3), Inches(3.2), Inches(4.2), Inches(3.5))
    tf = obs_box.text_frame
    p = tf.paragraphs[0]
    p.text = "关键观察"
    p.font.size = Pt(14)
    p.font.bold = True

    for obs in data_dict["observations"]:
        p = tf.add_paragraph()
        p.text = obs
        p.font.size = Pt(11)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # HLA_ESM2_2 analysis slide
    hla2_data = {
        "best_epoch": "7",
        "best_performance_avg": "0.8933",
        "no_improve_count": "4",
        "observations": [
            "• 收敛更快 (epoch 7 vs Phase1 epoch 10)",
            "• 最佳验证性能提升 +0.35%",
            "• Independent 测试集 ROC-AUC +0.19%",
            "• External 测试集 ROC-AUC +0.30%",
            "• encoder_P 迁移效果有效",
            "• 整体性能轻微提升",
        ],
    }

    hla2_test_sets = [
        ["Independent", "0.9641", "0.9123"],
        ["External", "0.9219", "0.8467"],
    ]

    hla2_comparison = [
        ["最佳 Epoch", "10", "7"],
        ["验证性能均值", "0.8898", "0.8933 (+0.35%)"],
        ["Ind. ROC-AUC", "0.9622", "0.9641 (+0.19%)"],
        ["Ind. Accuracy", "0.9103", "0.9123 (+0.20%)"],
        ["Ext. ROC-AUC", "0.9189", "0.9219 (+0.30%)"],
    ]

    add_phase2_slide(prs, "HLA_ESM2_2.py (Phase 2) 训练结果分析", hla2_data, hla2_test_sets, hla2_comparison)

    # TCR_ESM2_2 analysis slide
    tcr2_data = {
        "best_epoch": "7",
        "best_performance_avg": "0.8507",
        "no_improve_count": "4",
        "observations": [
            "• 收敛更快 (epoch 7 vs Phase1 epoch 8)",
            "• 最佳验证性能提升 +0.23%",
            "• Independent 测试集 ROC-AUC +0.23%",
            "• Triple 测试集 Accuracy +1.07%",
            "• ⚠️ Covid 测试集严重问题:",
            "   ROC-AUC 仅 0.5184 (下降 -3.57%)",
            "   Accuracy 仅 0.5211 (下降 -3.20%)",
            "• Covid 数据分布需深入分析",
        ],
    }

    tcr2_test_sets = [
        ["Independent", "0.9415", "0.8789"],
        ["Triple", "0.8349", "0.805"],
        ["Covid ⚠️", "0.5184", "0.5211"],
    ]

    tcr2_comparison = [
        ["最佳 Epoch", "8", "7"],
        ["验证性能均值", "0.8484", "0.8507 (+0.23%)"],
        ["Ind. ROC-AUC", "0.9392", "0.9415 (+0.23%)"],
        ["Ind. Accuracy", "0.8764", "0.8789 (+0.25%)"],
        ["Covid ROC-AUC", "0.5541", "0.5184 (-3.57%)"],
    ]

    add_phase2_slide(prs, "TCR_ESM2_2.py (Phase 2) 训练结果分析", tcr2_data, tcr2_test_sets, tcr2_comparison)

    # Save
    output_path = "/home/mclab/mjp/unifyimmun/docs/esm2_phase2_analysis.pptx"
    prs.save(output_path)
    print(f"PPT saved to: {output_path}")


if __name__ == "__main__":
    main()