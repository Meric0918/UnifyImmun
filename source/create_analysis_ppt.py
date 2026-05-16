"""
Generate analysis PPT for HLA_ESM2 and TCR_ESM2 training results.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


def add_title_slide(prs, title, subtitle):
    """Add a title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(9), Inches(0.5))
    tf = subtitle_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(18)
    p.alignment = PP_ALIGN.CENTER


def add_analysis_slide(prs, title, data_dict, test_sets_data=None):
    """Add an analysis slide with tables and metrics."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.3), Inches(9.4), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 51, 102)

    # Training summary table
    summary_box = slide.shapes.add_textbox(Inches(0.3), Inches(1), Inches(4.5), Inches(0.4))
    tf = summary_box.text_frame
    p = tf.paragraphs[0]
    p.text = "训练概况"
    p.font.size = Pt(16)
    p.font.bold = True

    # Create summary table
    rows = 4
    cols = 2
    table = slide.shapes.add_table(rows, cols, Inches(0.3), Inches(1.4), Inches(4.5), Inches(1.5)).table

    summary_data = [
        ["项目", "值"],
        ["最佳 Epoch", data_dict["best_epoch"]],
        ["最佳验证性能均值", data_dict["best_performance_avg"]],
        ["早停计数 (最后)", data_dict["no_improve_count"]],
    ]

    for i, row_data in enumerate(summary_data):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            cell.text = cell_text
            cell.text_frame.paragraphs[0].font.size = Pt(12)
            if i == 0:
                cell.text_frame.paragraphs[0].font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(200, 200, 200)

    # Validation metrics
    val_box = slide.shapes.add_textbox(Inches(0.3), Inches(3.1), Inches(4.5), Inches(0.4))
    tf = val_box.text_frame
    p = tf.paragraphs[0]
    p.text = "验证集关键指标 (最佳 Epoch)"
    p.font.size = Pt(16)
    p.font.bold = True

    val_rows = 6
    val_cols = 2
    val_table = slide.shapes.add_table(val_rows, val_cols, Inches(0.3), Inches(3.5), Inches(4.5), Inches(2)).table

    val_data = [
        ["指标", "值"],
        ["ROC-AUC", data_dict["val_roc_auc"]],
        ["Accuracy", data_dict["val_accuracy"]],
        ["MCC", data_dict["val_mcc"]],
        ["F1", data_dict["val_f1"]],
        ["Loss", data_dict["val_loss"]],
    ]

    for i, row_data in enumerate(val_data):
        for j, cell_text in enumerate(row_data):
            cell = val_table.cell(i, j)
            cell.text = cell_text
            cell.text_frame.paragraphs[0].font.size = Pt(12)
            if i == 0:
                cell.text_frame.paragraphs[0].font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(200, 200, 200)

    # Test sets performance (right side)
    test_box = slide.shapes.add_textbox(Inches(5), Inches(1), Inches(4.5), Inches(0.4))
    tf = test_box.text_frame
    p = tf.paragraphs[0]
    p.text = "测试集表现"
    p.font.size = Pt(16)
    p.font.bold = True

    if test_sets_data:
        test_rows = len(test_sets_data) + 1
        test_cols = 3
        test_table = slide.shapes.add_table(test_rows, test_cols, Inches(5), Inches(1.4), Inches(4.5), Inches(test_rows * 0.5)).table

        test_header = ["测试集", "ROC-AUC", "Accuracy"]
        for j, header in enumerate(test_header):
            cell = test_table.cell(0, j)
            cell.text = header
            cell.text_frame.paragraphs[0].font.size = Pt(12)
            cell.text_frame.paragraphs[0].font.bold = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(200, 200, 200)

        for i, row_data in enumerate(test_sets_data):
            for j, cell_text in enumerate(row_data):
                cell = test_table.cell(i + 1, j)
                cell.text = str(cell_text)
                cell.text_frame.paragraphs[0].font.size = Pt(12)
                # Highlight poor performance
                if j == 1 or j == 2:
                    try:
                        val = float(cell_text)
                        if val < 0.6:
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = RGBColor(255, 200, 200)
                        elif val > 0.9:
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = RGBColor(200, 255, 200)
                    except:
                        pass

    # Key observations
    obs_box = slide.shapes.add_textbox(Inches(5), Inches(4), Inches(4.5), Inches(3))
    tf = obs_box.text_frame
    p = tf.paragraphs[0]
    p.text = "关键观察"
    p.font.size = Pt(16)
    p.font.bold = True

    for obs in data_dict["observations"]:
        p = tf.add_paragraph()
        p.text = obs
        p.font.size = Pt(11)
        p.level = 0


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title slide
    add_title_slide(
        prs,
        "UnifyImmun ESM2 模型训练分析",
        "HLA_ESM2 与 TCR_ESM2 训练结果对比"
    )

    # HLA_ESM2 analysis slide
    hla_data = {
        "best_epoch": "10",
        "best_performance_avg": "0.8898",
        "no_improve_count": "4 (patience=5)",
        "val_roc_auc": "0.9624",
        "val_accuracy": "0.9109",
        "val_mcc": "0.7972",
        "val_f1": "0.864",
        "val_loss": "0.2909",
        "observations": [
            "• 早停未触发 (no_improve_count=4)",
            "• 最佳验证性能在 epoch 10 达到峰值",
            "• 过拟合迹象：val_loss 从 0.2362 升至 0.2909",
            "• 外部测试集稍弱 (acc 0.8561)",
            "• 整体性能优秀，泛化良好",
        ],
    }

    hla_test_sets = [
        ["Independent", "0.9622", "0.9103"],
        ["External", "0.9189", "0.8561"],
    ]

    add_analysis_slide(prs, "HLA_ESM2.py 训练结果分析", hla_data, hla_test_sets)

    # TCR_ESM2 analysis slide
    tcr_data = {
        "best_epoch": "8",
        "best_performance_avg": "0.8484",
        "no_improve_count": "4 (patience=5)",
        "val_roc_auc": "0.9372",
        "val_accuracy": "0.8761",
        "val_mcc": "0.7238",
        "val_f1": "0.8169",
        "val_loss": "0.386",
        "observations": [
            "• 训练收敛更快 (最佳 epoch 8)",
            "• encoder_P 迁移有效",
            "• 过拟合比 HLA 更严重",
            "• Covid 测试集严重问题 ⚠️",
            "• ROC-AUC 仅 0.5541 (接近随机)",
            "• 需深入分析 Covid 数据分布",
        ],
    }

    tcr_test_sets = [
        ["Independent", "0.9392", "0.8764"],
        ["Triple", "0.8485", "0.7943"],
        ["Covid ⚠️", "0.5541", "0.5531"],
    ]

    add_analysis_slide(prs, "TCR_ESM2.py 训练结果分析", tcr_data, tcr_test_sets)

    # Save
    output_path = "/home/mclab/mjp/unifyimmun/docs/esm2_training_analysis.pptx"
    prs.save(output_path)
    print(f"PPT saved to: {output_path}")


if __name__ == "__main__":
    main()